"""Secure in-memory parsing of validated uploads.

Parsing operates on the uploaded bytes in memory: PyMuPDF opens the PDF from a
byte stream and python-docx opens the DOCX from a :class:`io.BytesIO`. Only
text is extracted; macros, embedded scripts, and OLE objects are never executed
or evaluated. Every failure mode (encrypted, malformed, empty file, empty
extraction) raises a handled :class:`DocumentError` mapped to a friendly
message rather than crashing. The upload is never mutated or written back, and
any unavoidable temp artifact is removed in a ``finally`` block.
"""

from __future__ import annotations

from io import BytesIO
from typing import TYPE_CHECKING

from hirely.schemas.documents import ParsedDocument

if TYPE_CHECKING:
    from hirely.documents.validation import UploadLike


class DocumentError(Exception):
    """A validated upload could not be parsed into text.

    The message is user-safe and names the reason (encrypted, corrupt, empty,
    or no extractable text) without leaking document content or a stack trace.
    """


def _parse_pdf(data: bytes) -> tuple[str, int]:
    """Extract text from PDF bytes in memory.

    Args:
        data: The validated PDF bytes.

    Returns:
        A tuple of the extracted text and the page count.

    Raises:
        DocumentError: If the PDF is encrypted or malformed.
    """
    import fitz  # PyMuPDF; imported lazily so module import stays dependency-light.

    document = None
    try:
        try:
            document = fitz.open(stream=data, filetype="pdf")
        except Exception as exc:  # noqa: BLE001 - PyMuPDF raises library-specific types.
            raise DocumentError("The PDF could not be read. It may be corrupted.") from exc

        if document.needs_pass:
            raise DocumentError("The PDF is password protected. Remove the password and upload again.")

        parts = [page.get_text() for page in document]
        page_count = document.page_count
        return "\n".join(parts), page_count
    finally:
        if document is not None:
            document.close()


def _parse_docx(data: bytes) -> str:
    """Extract text from DOCX bytes in memory.

    Args:
        data: The validated DOCX bytes.

    Returns:
        The extracted text.

    Raises:
        DocumentError: If the DOCX is malformed.
    """
    from docx import Document  # python-docx; imported lazily.

    try:
        document = Document(BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - python-docx raises library-specific types.
        raise DocumentError("The Word document could not be read. It may be corrupted.") from exc

    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.append(cell.text)
    return "\n".join(paragraphs)


def parse_document(upload: UploadLike, media_type: str) -> ParsedDocument:
    """Parse a validated upload into a :class:`ParsedDocument`.

    Call :func:`hirely.documents.validation.validate_upload` first; this
    function trusts the resolved ``media_type`` and parses the bytes in memory.

    Args:
        upload: The Streamlit upload (matching :class:`UploadLike`).
        media_type: The confirmed media type, ``"pdf"`` or ``"docx"``.

    Returns:
        The extracted text wrapped in a :class:`ParsedDocument`.

    Raises:
        DocumentError: If the file is empty, encrypted, malformed, or yields no
            extractable text.
    """
    data = upload.getvalue()
    if len(data) == 0:
        raise DocumentError("The uploaded file is empty.")

    page_count: int | None = None
    if media_type == "pdf":
        text, page_count = _parse_pdf(data)
    elif media_type == "docx":
        text = _parse_docx(data)
    else:  # Defensive: validation only ever yields "pdf" or "docx".
        raise DocumentError("Unsupported document type.")

    text = text.strip()
    if not text:
        raise DocumentError("No text could be extracted from the document. It may be a scanned image.")

    return ParsedDocument(
        filename=upload.name,
        media_type=media_type,  # type: ignore[arg-type]
        text=text,
        page_count=page_count,
        char_count=len(text),
    )
